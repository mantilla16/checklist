"""Genera el libro "Revisión Pagos a Proveedores" a partir de una plantilla.

El .xlsx de la carpeta del proyecto se usa SOLO como plantilla: de ahi se toma
el formato (estilos, encabezado del lote, fila de titulos, filas de Total y
Diferencias). Los datos los captura el usuario en la interfaz.

Estructura que se reproduce en cada hoja de lote:

- B2:C8  encabezado del lote (tipo de pago, nombre, cuenta, valor total...).
- Fila de titulos con "Nro. Registro" en la columna B.
- Un registro por fila. Si el pago se reparte en VARIAS FACTURAS el registro
  ocupa varias filas: cada una con su K, y la columna L combinada con la
  formula =+G12-K12-K13 (una resta por sub-fila).
- Fila "Total" con =SUM(...) y fila "Diferencias" con =+G17-C5.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

COL = {
    "nro": 2, "estado": 3, "titular": 4, "documento": 5, "cuenta": 6, "valor": 7,
    "correo": 10, "valor_aprobado": 11, "diferencias": 12,
    "factura": 13, "no_factura": 14, "nit_tercero": 15, "validacion": 16,
    "valor_pagar": 17, "dif_pago": 18, "radian": 19, "orden_compra": 20,
    "egreso": 21, "compras": 22, "entrada": 23, "recibido": 24,
    "comentarios": 25,
}
# Columnas que se combinan cuando el registro ocupa varias filas: son del
# registro completo, no de cada factura (asi esta en el archivo original).
COMBINADAS = ("diferencias", "factura", "validacion", "dif_pago", "comentarios")
# Columnas que van una por factura
POR_RENGLON = ("valor_aprobado", "no_factura", "nit_tercero", "valor_pagar",
               "radian", "orden_compra", "egreso", "compras", "entrada",
               "recibido")
PRIMERA_COL, ULTIMA_COL = 2, 25  # B..Y

ETIQUETAS_INFO = {
    "tipo de pago": "tipo_pago",
    "nombre del pago": "nombre_pago",
    "cuenta a debitar": "cuenta",
    "valor total": "valor_total",
    "numero total de registros": "num_registros",
    "fecha de creacion del lote": "fecha_creacion",
    "fecha de aplicacion": "fecha_aplicacion",
}


def _limpiar(texto) -> str:
    """Normaliza etiquetas: quita espacio duro, dos puntos y tildes."""
    if texto is None:
        return ""
    base = str(texto).replace("\xa0", " ").strip().rstrip(":").strip().lower()
    for antes, despues in {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u", "ñ": "n"}.items():
        base = base.replace(antes, despues)
    return base


# --------------------------------------------------------------------------- #
# Lectura de la plantilla (solo el formato)
# --------------------------------------------------------------------------- #

def _disposicion(hoja) -> dict | None:
    """Ubica las filas clave de una hoja de lote. None si no es una."""
    fila_titulos = None
    for fila in range(1, min(hoja.max_row, 40) + 1):
        if _limpiar(hoja.cell(row=fila, column=COL["nro"]).value) == "nro. registro":
            fila_titulos = fila
            break
    if fila_titulos is None:
        return None

    etiquetas = {}
    for fila in range(1, fila_titulos):
        clave = ETIQUETAS_INFO.get(_limpiar(hoja.cell(row=fila, column=2).value))
        if clave:
            etiquetas[clave] = fila

    fila_total = None
    fila_diferencias = None
    for fila in range(fila_titulos + 1, hoja.max_row + 2):
        etiqueta = _limpiar(hoja.cell(row=fila, column=COL["cuenta"]).value)
        if etiqueta == "total" and fila_total is None:
            fila_total = fila
        elif etiqueta == "diferencias":
            fila_diferencias = fila

    return {
        "fila_titulos": fila_titulos,
        "fila_total": fila_total or hoja.max_row + 1,
        "fila_diferencias": fila_diferencias,
        "filas_etiquetas": etiquetas,
        "capacidad": (fila_total or hoja.max_row + 1) - fila_titulos - 1,
    }


def leer_plantilla(ruta: str | Path) -> dict:
    """Describe la plantilla: que hoja sirve de modelo y con que disposicion."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    modelo = None
    for hoja in libro.worksheets:
        disp = _disposicion(hoja)
        if disp:
            modelo = {"hoja": hoja.title, **disp}
            break

    if modelo is None:
        return {"valida": False,
                "error": "La plantilla no tiene una hoja con la tabla de registros."}

    hoja = libro[modelo["hoja"]]
    titulos = {
        get_column_letter(c): hoja.cell(row=modelo["fila_titulos"], column=c).value
        for c in range(PRIMERA_COL, ULTIMA_COL + 1)
        if hoja.cell(row=modelo["fila_titulos"], column=c).value
    }
    ejemplo = {
        clave: hoja.cell(row=fila, column=3).value
        for clave, fila in modelo["filas_etiquetas"].items()
    }
    return {
        "valida": True,
        "modelo": modelo,
        "titulos": titulos,
        "ejemplo_info": {
            k: (v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else v)
            for k, v in ejemplo.items()
        },
        "hojas_lote": [h.title for h in libro.worksheets if _disposicion(h)],
    }


# --------------------------------------------------------------------------- #
# Escritura
# --------------------------------------------------------------------------- #

# Observacion de cada renglon -> columna a la que pertenece. La Y las reune
# todas para poder leer de un tiron que le falta al registro.
OBSERVACIONES = (
    ("obs_no_factura", "N"), ("obs_nit_tercero", "O"), ("obs_valor_pagar", "Q"),
    ("obs_radian", "S"), ("obs_orden_compra", "T"), ("obs_egreso_g", "U"),
    ("obs_compras_p", "V"), ("obs_entrada_e", "W"), ("obs_recibido_x", "X"),
)
# Columnas cuyo estado se anota aunque no traigan observacion escrita
ESTADOS = (
    ("radian", "S"), ("orden_compra", "T"), ("egreso_g", "U"),
    ("compras_p", "V"), ("entrada_e", "W"), ("recibido_x", "X"),
)

ETIQUETA_FALTANTE = {
    "nit_cliente": "NIT del cliente",
    "nit_tercero": "NIT del tercero",
    "cufe": "CUFE",
    "qr": "código QR",
    "numero": "número de factura",
}


def _comentarios(renglones: list[dict], nota_factura: str = "") -> str:
    """Reune en un texto lo que dijo cada paso, con la columna de la que sale.

    Cada observacion sigue estando en su propia celda; esta columna es el
    resumen, para no tener que recorrer catorce columnas por registro.
    """
    lineas: list[str] = []
    if nota_factura:
        lineas.append(f"M: {nota_factura}")

    varios = len(renglones) > 1
    for renglon in renglones:
        prefijo = ""
        if varios:
            numero = renglon.get("no_factura") or renglon.get("factura") or ""
            prefijo = f"{numero} · " if numero else ""

        for clave, columna in OBSERVACIONES:
            texto = (renglon.get(clave) or "").strip()
            if texto:
                lineas.append(f"{prefijo}{columna}: {texto}")

        # Un "Pendiente" sin observacion tambien es informacion: dice que ese
        # control no se pudo hacer, no que este bien
        for clave, columna in ESTADOS:
            estado = renglon.get(clave)
            observado = any(renglon.get(o) for o, c in OBSERVACIONES if c == columna)
            if estado == "Pendiente" and not observado:
                lineas.append(f"{prefijo}{columna}: pendiente")

    # Sin repetir: dos facturas del mismo registro suelen fallar por lo mismo
    return "\n".join(dict.fromkeys(lineas))


def _escribir_con_nota(hoja, fila: int, columna: int, valor, nota: str | None) -> None:
    """Escribe el valor y, si hay observacion, la deja en la MISMA celda.

    La observacion de un hallazgo va en la columna a la que corresponde: si es
    del numero de factura, en N; si es del NIT del tercero, en O. Queda debajo
    del valor dentro de la celda, con ajuste de texto.
    """
    celda = hoja.cell(row=fila, column=columna)
    if valor in (None, "") and not nota:
        celda.value = None
        return

    celda.value = f"{valor}\n{nota}" if nota else valor
    if nota:
        celda.alignment = Alignment(
            wrap_text=True,
            horizontal=celda.alignment.horizontal,
            vertical="top",
        )


def _limpiar_hoja(hoja, disp: dict) -> None:
    """Deja la copia de la plantilla en blanco: sin datos y sin combinaciones."""
    primera, fin = disp["fila_titulos"] + 1, disp["fila_total"] - 1
    for rango in list(hoja.merged_cells.ranges):
        if rango.min_row >= primera and rango.max_row <= fin:
            hoja.unmerge_cells(str(rango))
    for fila in range(primera, fin + 1):
        for col in range(PRIMERA_COL, ULTIMA_COL + 1):
            hoja.cell(row=fila, column=col).value = None
    for fila in disp["filas_etiquetas"].values():
        hoja.cell(row=fila, column=3).value = None


def _escribir_lote(hoja, disp: dict, lote: dict) -> list[dict]:
    """Escribe el encabezado y los registros de un lote. Devuelve el resumen."""
    info = dict(lote.get("info") or {})
    registros = lote.get("registros") or []
    # El numero de registros lo declara el lote del banco; si no viene, se
    # escribe el conteo de lo capturado.
    if info.get("num_registros") in (None, ""):
        info["num_registros"] = len(registros)

    for clave, fila in disp["filas_etiquetas"].items():
        valor = info.get(clave)
        if valor not in (None, ""):
            hoja.cell(row=fila, column=3).value = valor

    primera = disp["fila_titulos"] + 1
    fila_total = disp["fila_total"]

    necesarias = sum(max(1, len(r.get("renglones") or [])) for r in registros)
    if necesarias > fila_total - primera:
        faltan = necesarias - (fila_total - primera)
        hoja.insert_rows(fila_total, faltan)
        fila_total += faltan

    resumen = []
    fila = primera
    for reg in registros:
        renglones = reg.get("renglones") or [{}]
        bloque = max(1, len(renglones))

        hoja.cell(row=fila, column=COL["nro"]).value = reg.get("nro")
        hoja.cell(row=fila, column=COL["estado"]).value = reg.get("estado") or "Revisión"
        hoja.cell(row=fila, column=COL["titular"]).value = reg.get("titular")
        hoja.cell(row=fila, column=COL["documento"]).value = reg.get("documento")
        hoja.cell(row=fila, column=COL["cuenta"]).value = reg.get("cuenta")
        hoja.cell(row=fila, column=COL["valor"]).value = reg.get("valor")

        validado = any(r.get("valor") is not None for r in renglones)
        for i, renglon in enumerate(renglones):
            hoja.cell(row=fila + i, column=COL["correo"]).value = reg.get("correo") or None
            hoja.cell(row=fila + i, column=COL["valor_aprobado"]).value = renglon.get("valor")
            # N y O: una por factura, de la validacion de la factura electronica.
            # La observacion viaja en la celda de SU columna, debajo del valor.
            _escribir_con_nota(
                hoja, fila + i, COL["no_factura"],
                renglon.get("no_factura") or renglon.get("factura"),
                renglon.get("obs_no_factura"),
            )
            _escribir_con_nota(
                hoja, fila + i, COL["nit_tercero"],
                renglon.get("nit_tercero"), renglon.get("obs_nit_tercero"),
            )
            # Q: valor a pagar, del total del registro contable P
            _escribir_con_nota(
                hoja, fila + i, COL["valor_pagar"],
                renglon.get("valor_pagar"), renglon.get("obs_valor_pagar"),
            )
            # S: aceptacion en RADIAN, una por factura
            _escribir_con_nota(
                hoja, fila + i, COL["radian"],
                renglon.get("radian"), renglon.get("obs_radian"),
            )
            # T: orden de compra (Y)
            _escribir_con_nota(
                hoja, fila + i, COL["orden_compra"],
                renglon.get("orden_compra"), renglon.get("obs_orden_compra"),
            )
            # U: egreso (G) contra el valor de la tabla azul
            _escribir_con_nota(
                hoja, fila + i, COL["egreso"],
                renglon.get("egreso_g"), renglon.get("obs_egreso_g"),
            )
            # V: registro contable de compras (P)
            _escribir_con_nota(
                hoja, fila + i, COL["compras"],
                renglon.get("compras_p"), renglon.get("obs_compras_p"),
            )
            # W: entrada de inventario (E)
            _escribir_con_nota(
                hoja, fila + i, COL["entrada"],
                renglon.get("entrada_e"), renglon.get("obs_entrada_e"),
            )
            # X: relacion de mercancia recibida (F-CO-110)
            _escribir_con_nota(
                hoja, fila + i, COL["recibido"],
                renglon.get("recibido_x"), renglon.get("obs_recibido_x"),
            )

        # L: una celda (combinada si hay varias facturas) con la resta por cada
        # sub-fila. Va SIEMPRE: es una formula, y con la K vacia calcula G - 0,
        # o sea el valor que todavia falta por conciliar.
        restas = "".join(f"-K{fila + i}" for i in range(bloque))
        hoja.cell(row=fila, column=COL["diferencias"]).value = f"=+G{fila}{restas}"

        # R: la misma resta pero contra el valor a pagar (Q)
        restas_q = "".join(f"-Q{fila + i}" for i in range(bloque))
        hoja.cell(row=fila, column=COL["dif_pago"]).value = f"=+G{fila}{restas_q}"

        # M y P: del registro completo. M es "OK" si todas sus facturas pasaron
        # las revisiones; P es VERDADERO si todos los NIT tercero coinciden.
        # La condicion es que la factura se haya cargado, no que su validacion
        # tenga valor: la P puede quedar sin determinar cuando falta la tabla
        # azul, y con la condicion anterior la M dejaba de escribirse justo en
        # el caso en que hay algo que decir.
        notas: list[str] = []
        con_factura = [r for r in renglones if r.get("archivo_factura")
                       or r.get("faltantes") or r.get("pendientes")]
        if con_factura:
            todas_ok = all(r.get("factura_ok") for r in renglones)
            faltantes = sorted({
                falta for r in renglones for falta in (r.get("faltantes") or [])
            })
            pendientes = sorted({
                falta for r in renglones for falta in (r.get("pendientes") or [])
            } - set(faltantes))
            etiqueta = lambda claves: ", ".join(
                ETIQUETA_FALTANTE.get(c, c) for c in claves)
            if faltantes:
                notas.append("falta: " + etiqueta(faltantes))
            # Lo que no se pudo comprobar se dice, en vez de callarlo y dejar
            # que la celda parezca revisada
            if pendientes:
                notas.append("sin comprobar: " + etiqueta(pendientes))
            _escribir_con_nota(
                hoja, fila, COL["factura"],
                "OK" if todas_ok else "Revisar",
                " · ".join(notas) if notas else None,
            )
            # P: VERDADERO solo si TODOS los NIT tercero coinciden. Si alguno
            # quedo sin determinar (falta la tabla azul) la celda va vacia: un
            # FALSO ahi seria una acusacion sin haber comparado nada.
            validaciones = [r.get("validacion") for r in renglones]
            hoja.cell(row=fila, column=COL["validacion"]).value = (
                None if any(v is None for v in validaciones) else all(validaciones))

        # Y: los comentarios de todos los pasos, reunidos
        comentarios = _comentarios(renglones, " · ".join(notas))
        celda_y = hoja.cell(row=fila, column=COL["comentarios"])
        celda_y.value = comentarios or None
        if comentarios:
            celda_y.alignment = Alignment(wrap_text=True, vertical="top",
                                          horizontal=celda_y.alignment.horizontal)

        for clave in COMBINADAS:
            if bloque > 1:
                letra = get_column_letter(COL[clave])
                hoja.merge_cells(f"{letra}{fila}:{letra}{fila + bloque - 1}")

        resumen.append({
            "hoja": hoja.title, "registro": reg.get("nro"), "fila": fila,
            "facturas": bloque, "validado": validado,
            "titular": reg.get("titular"),
            "valor_lote": reg.get("valor") or 0,
            "valor_aprobado": sum(r.get("valor") or 0 for r in renglones),
        })
        fila += bloque

    ultima = max(fila - 1, primera)
    g, k, ele = (get_column_letter(COL[c]) for c in ("valor", "valor_aprobado", "diferencias"))
    hoja.cell(row=fila_total, column=COL["cuenta"]).value = "Total"
    hoja.cell(row=fila_total, column=COL["valor"]).value = f"=SUM({g}{primera}:{g}{ultima})"
    hoja.cell(row=fila_total, column=COL["valor_aprobado"]).value = f"=SUM({k}{primera}:{k}{ultima})"
    hoja.cell(row=fila_total, column=COL["diferencias"]).value = f"=SUM({ele}{primera}:{ele}{ultima})"
    q, erre = (get_column_letter(COL[c]) for c in ("valor_pagar", "dif_pago"))
    hoja.cell(row=fila_total, column=COL["valor_pagar"]).value = f"=SUM({q}{primera}:{q}{ultima})"
    hoja.cell(row=fila_total, column=COL["dif_pago"]).value = f"=SUM({erre}{primera}:{erre}{ultima})"

    # La fila "Diferencias" compara el total sumado contra el valor del lote (C5)
    fila_dif = disp["fila_diferencias"]
    if fila_dif:
        fila_dif += fila_total - disp["fila_total"]
        fila_valor_total = disp["filas_etiquetas"].get("valor_total")
        referencia = f"C{fila_valor_total}" if fila_valor_total else "C5"
        hoja.cell(row=fila_dif, column=COL["cuenta"]).value = "Diferencias"
        hoja.cell(row=fila_dif, column=COL["valor"]).value = f"=+{g}{fila_total}-{referencia}"

    return resumen


def exportar(ruta_plantilla: str | Path, ruta_destino: str | Path,
             lotes: list[dict]) -> dict:
    """Genera el libro final: una hoja por lote, con el formato de la plantilla."""
    descripcion = leer_plantilla(ruta_plantilla)
    if not descripcion["valida"]:
        raise ValueError(descripcion["error"])
    if not lotes:
        raise ValueError("No hay lotes para exportar")

    modelo = descripcion["modelo"]
    disp = {clave: valor for clave, valor in modelo.items() if clave != "hoja"}
    libro = openpyxl.load_workbook(ruta_plantilla)   # con formulas y estilos
    hoja_modelo = libro[modelo["hoja"]]

    # Las hojas de lote de la plantilla son un ejemplo y se van a borrar. Se
    # renombran de una vez para que el usuario pueda llamar su lote igual que
    # la plantilla sin que el nombre choque.
    temporales = {}
    for titulo in descripcion["hojas_lote"]:
        temporales[titulo] = f"~plantilla {len(temporales)}"
        libro[titulo].title = temporales[titulo]

    resumen, creadas = [], []
    for i, lote in enumerate(lotes):
        nombre = str(lote.get("hoja") or f"Lote {i + 1}").strip()[:31] or f"Lote {i + 1}"
        if nombre in libro.sheetnames:
            nombre = f"{nombre[:27]} ({i + 1})"

        hoja = libro.copy_worksheet(hoja_modelo)
        hoja.title = nombre
        _limpiar_hoja(hoja, disp)
        resumen.extend(_escribir_lote(hoja, dict(disp), lote))
        creadas.append(nombre)

    # Fuera las hojas de ejemplo de la plantilla
    for titulo in temporales.values():
        if titulo in libro.sheetnames:
            del libro[titulo]

    libro.save(ruta_destino)
    return {"archivo": str(ruta_destino), "hojas": creadas, "registros": resumen}
